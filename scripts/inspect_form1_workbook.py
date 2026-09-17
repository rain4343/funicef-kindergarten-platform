from pathlib import Path
from openpyxl import load_workbook

path = Path('/home/ubuntu/upload/Form1.xlsx')
wb = load_workbook(path, data_only=False)
print('SHEETS', wb.sheetnames)
for ws in wb.worksheets:
    print(f'\nSHEET {ws.title} dimensions={ws.max_row}x{ws.max_column}')
    print('MERGED', list(ws.merged_cells.ranges))
    print('FREEZE', ws.freeze_panes)
    print('FILTER', ws.auto_filter.ref)
    for row in ws.iter_rows():
        values = [cell.value for cell in row]
        if any(v is not None for v in values):
            print('ROW', row[0].row, values)
    print('COLUMN_WIDTHS', {k: v.width for k, v in ws.column_dimensions.items() if v.width})
    print('ROW_HEIGHTS', {k: v.height for k, v in ws.row_dimensions.items() if v.height})
    print('PRINT', ws.print_area, ws.page_setup.orientation, ws.page_setup.paperSize)
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                print('STYLE', cell.coordinate, 'value=',repr(cell.value), 'font=',cell.font.name,cell.font.sz,cell.font.bold,cell.font.italic,cell.font.color.type if cell.font.color else None,cell.font.color.rgb if cell.font.color and cell.font.color.type=='rgb' else None, 'fill=',cell.fill.fill_type,cell.fill.fgColor.type,cell.fill.fgColor.rgb, 'align=',cell.alignment.horizontal,cell.alignment.vertical,cell.alignment.wrap_text, 'border=',cell.border.left.style,cell.border.right.style,cell.border.top.style,cell.border.bottom.style, 'numfmt=',cell.number_format)
